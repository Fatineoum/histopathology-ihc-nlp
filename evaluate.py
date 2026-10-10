"""
evaluate.py
IHC NLP Project : Evaluation against Ground Truth
--------------------------------------------------
Runs the extraction pipeline against reports_ground_truth.xlsx and prints
precision / recall / F1 for positive and negative IHC marker extraction.

Usage:
    python evaluate.py                        # rule-based pipeline only
    python evaluate.py --hybrid               # hybrid pipeline (requires trained model)
    python evaluate.py --by-diagnosis         # per-diagnosis F1 breakdown
    python evaluate.py --output results/eval.csv  # save per-case results

Ground truth columns (row 2 = headers, data from row 3):
    col  0 : Patient ID
    col  4 : Final Diagnosis
    col 17 : IHC text
    col 18 : Positive IHC (ground truth)
    col 19 : Negative IHC (ground truth)
"""

import argparse
import sys
from pathlib import Path
from collections import defaultdict

import pandas as pd
import openpyxl

DATA_DIR = Path("data")
GT_FILE  = DATA_DIR / "reports_ground_truth.xlsx"

# Column indices (0-based) in the ground truth sheet
COL_PID      = 0
COL_DX       = 4
COL_IHC_TEXT = 17
COL_POS_GT   = 18
COL_NEG_GT   = 19


# ── Load ground truth ─────────────────────────────────────────────────────────

def load_ground_truth():
    print(f"Loading ground truth from {GT_FILE} ...", flush=True)
    wb = openpyxl.load_workbook(GT_FILE, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # Row 0 = merged header, row 1 = actual column headers, data from row 2
    data = []
    for row in rows[2:]:
        pid      = row[COL_PID]
        ihc_text = row[COL_IHC_TEXT]
        pos_gt   = row[COL_POS_GT]
        neg_gt   = row[COL_NEG_GT]
        dx       = row[COL_DX]
        if not pid or (not ihc_text and not pos_gt and not neg_gt):
            continue
        data.append({
            "patient_id": str(pid).strip(),
            "diagnosis":  str(dx).strip().lower() if dx else "unknown",
            "ihc_text":   str(ihc_text).strip() if ihc_text else "",
            "pos_gt":     str(pos_gt).strip()   if pos_gt   else "",
            "neg_gt":     str(neg_gt).strip()   if neg_gt   else "",
        })
    wb.close()
    df = pd.DataFrame(data)
    # Only keep cases with at least one ground truth label
    df = df[(df["pos_gt"].str.lower().isin({"", "nil", "na", "n/a"}) == False)
            | (df["neg_gt"].str.lower().isin({"", "nil", "na", "n/a"}) == False)]
    print(f"  {len(df)} cases with ground truth labels found.\n", flush=True)
    return df


# ── Borrow helpers from 03_rule_based_pipeline ───────────────────────────────

def load_rule_module():
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "rule_pipeline", Path("03_rule_based_pipeline.py"))
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# ── Metric helpers ────────────────────────────────────────────────────────────

def prf(tp, fp, fn):
    p = tp / (tp + fp) if (tp + fp) > 0 else 0.0
    r = tp / (tp + fn) if (tp + fn) > 0 else 0.0
    f = 2 * p * r / (p + r) if (p + r) > 0 else 0.0
    return round(p, 4), round(r, 4), round(f, 4)


def print_results_table(label, totals):
    tp_p = totals["tp_pos"]; fp_p = totals["fp_pos"]; fn_p = totals["fn_pos"]
    tp_n = totals["tp_neg"]; fp_n = totals["fp_neg"]; fn_n = totals["fn_neg"]
    p_p, r_p, f_p = prf(tp_p, fp_p, fn_p)
    p_n, r_n, f_n = prf(tp_n, fp_n, fn_n)

    # Micro-average across pos + neg
    tp_all = tp_p + tp_n; fp_all = fp_p + fp_n; fn_all = fn_p + fn_n
    p_all, r_all, f_all = prf(tp_all, fp_all, fn_all)

    col = 14
    print(f"\n{'=' * 60}")
    print(f"  {label}")
    print(f"{'=' * 60}")
    print(f"  {'':20s}  {'Precision':>10}  {'Recall':>8}  {'F1':>8}")
    print(f"  {'-' * 52}")
    print(f"  {'Positive IHC':20s}  {p_p:>10.1%}  {r_p:>8.1%}  {f_p:>8.1%}   "
          f"(TP={tp_p} FP={fp_p} FN={fn_p})")
    print(f"  {'Negative IHC':20s}  {p_n:>10.1%}  {r_n:>8.1%}  {f_n:>8.1%}   "
          f"(TP={tp_n} FP={fp_n} FN={fn_n})")
    print(f"  {'-' * 52}")
    print(f"  {'Micro-average':20s}  {p_all:>10.1%}  {r_all:>8.1%}  {f_all:>8.1%}")
    print(f"{'=' * 60}\n")

    return {
        "pipeline":       label,
        "pos_precision":  p_p, "pos_recall": r_p, "pos_f1": f_p,
        "neg_precision":  p_n, "neg_recall": r_n, "neg_f1": f_n,
        "micro_precision": p_all, "micro_recall": r_all, "micro_f1": f_all,
    }


# ── Rule-based evaluation ─────────────────────────────────────────────────────

def run_rule_eval(df, mod):
    totals = defaultdict(int)
    rows = []
    for _, row in df.iterrows():
        preds   = mod.extract_from_text(row["ihc_text"])
        gt_pos  = mod.parse_gt_list(row["pos_gt"])
        gt_neg  = mod.parse_gt_list(row["neg_gt"])
        stats   = mod.evaluate_case(preds, gt_pos, gt_neg)
        for k, v in stats.items():
            totals[k] += v
        rows.append({
            "patient_id": row["patient_id"],
            "tp_pos": stats["tp_pos"], "fp_pos": stats["fp_pos"], "fn_pos": stats["fn_pos"],
            "tp_neg": stats["tp_neg"], "fp_neg": stats["fp_neg"], "fn_neg": stats["fn_neg"],
        })
    return totals, pd.DataFrame(rows)


# ── Per-diagnosis breakdown ───────────────────────────────────────────────────

def print_by_diagnosis(case_df, label, min_cases=10):
    """Print micro-averaged F1 per diagnosis, sorted descending. Only diagnoses
    with at least `min_cases` in the ground truth are shown."""
    print(f"\n{'=' * 70}")
    print(f"  {label} — per-diagnosis breakdown (min {min_cases} cases)")
    print(f"{'=' * 70}")
    print(f"  {'Diagnosis':<40s}  {'Cases':>5}  {'Pos F1':>7}  {'Neg F1':>7}  {'Overall F1':>10}")
    print(f"  {'-' * 66}")

    case_df = case_df[~case_df["diagnosis"].str.contains("unconfirmed|unconmfirmed", regex=True)]
    dx_groups = case_df.groupby("diagnosis")
    rows = []
    for dx, grp in dx_groups:
        if len(grp) < min_cases:
            continue
        t = grp[["tp_pos","fp_pos","fn_pos","tp_neg","fp_neg","fn_neg"]].sum()
        _, _, f_p = prf(t["tp_pos"], t["fp_pos"], t["fn_pos"])
        _, _, f_n = prf(t["tp_neg"], t["fp_neg"], t["fn_neg"])
        tp_all = t["tp_pos"] + t["tp_neg"]
        fp_all = t["fp_pos"] + t["fp_neg"]
        fn_all = t["fn_pos"] + t["fn_neg"]
        _, _, f_all = prf(tp_all, fp_all, fn_all)
        rows.append((dx, len(grp), f_p, f_n, f_all))

    rows.sort(key=lambda x: x[4], reverse=True)
    for dx, n, f_p, f_n, f_all in rows:
        label_str = dx[:38] + ".." if len(dx) > 40 else dx
        print(f"  {label_str:<40s}  {n:>5}  {f_p:>7.3f}  {f_n:>7.3f}  {f_all:>10.3f}")

    if rows:
        f1s = [r[4] for r in rows]
        ns  = [r[1] for r in rows]
        wmean = sum(f * n for f, n in zip(f1s, ns)) / sum(ns)
        median = sorted(f1s)[len(f1s) // 2]
        print(f"  {'-' * 66}")
        print(f"  Weighted mean F1: {wmean:.3f}   Median F1: {median:.3f}   "
              f"Diagnoses shown: {len(rows)}")
    print(f"{'=' * 70}\n")


# ── Hybrid evaluation ─────────────────────────────────────────────────────────

def run_hybrid_eval(df, mod):
    try:
        import importlib.util
        spec = importlib.util.spec_from_file_location(
            "hybrid_pipeline", Path("05_hybrid_pipeline.py"))
        hyb = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(hyb)
    except Exception as e:
        print(f"Could not load hybrid pipeline: {e}")
        sys.exit(1)

    rule_extract = hyb.load_rule_pipeline()
    model, tokenizer, predict_fn = hyb.load_bert_pipeline()

    totals = defaultdict(int)
    rows = []
    for _, row in df.iterrows():
        pred_pos, pred_neg, _ = hyb.combine(
            rule_extract(row["ihc_text"]), *predict_fn(row["ihc_text"], model, tokenizer)
        )
        preds_fmt = (
            [{"marker": m, "outcome": "positive"} for m in pred_pos] +
            [{"marker": m, "outcome": "negative"} for m in pred_neg]
        )
        gt_pos = mod.parse_gt_list(row["pos_gt"])
        gt_neg = mod.parse_gt_list(row["neg_gt"])
        stats  = mod.evaluate_case(preds_fmt, gt_pos, gt_neg)
        for k, v in stats.items():
            totals[k] += v
        rows.append({"patient_id": row["patient_id"], **stats})
    return totals, pd.DataFrame(rows)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Evaluate IHC extraction pipeline against ground truth")
    parser.add_argument("--hybrid",       action="store_true",
                        help="Also run hybrid (BERT + rule) pipeline")
    parser.add_argument("--by-diagnosis", action="store_true",
                        help="Print per-diagnosis F1 breakdown (min 10 cases per diagnosis)")
    parser.add_argument("--min-cases",    type=int, default=10,
                        help="Minimum cases for a diagnosis to appear in --by-diagnosis (default 10)")
    parser.add_argument("--output",       type=str, default=None,
                        help="Path to save per-case CSV results")
    args = parser.parse_args()

    df  = load_ground_truth()
    mod = load_rule_module()

    summary_rows = []

    print("Running rule-based pipeline ...", flush=True)
    rule_totals, rule_df = run_rule_eval(df, mod)
    summary_rows.append(
        print_results_table("Rule-based pipeline", rule_totals))

    if args.by_diagnosis:
        print_by_diagnosis(rule_df.merge(df[["patient_id","diagnosis"]], on="patient_id", how="left"),
                           "Rule-based pipeline", min_cases=args.min_cases)

    if args.hybrid:
        print("Running hybrid pipeline ...", flush=True)
        hyb_totals, hyb_df = run_hybrid_eval(df, mod)
        summary_rows.append(
            print_results_table("Hybrid pipeline (BERT-primary)", hyb_totals))
        if args.by_diagnosis:
            print_by_diagnosis(hyb_df.merge(df[["patient_id","diagnosis"]], on="patient_id", how="left"),
                               "Hybrid pipeline", min_cases=args.min_cases)

    if args.output:
        out = Path(args.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        rule_df.to_csv(out, index=False)
        print(f"Per-case results saved to {out}")

    # Summary table for README
    print("\nSummary (copy into README):")
    print(f"| Pipeline | Pos Precision | Pos Recall | Pos F1 | Neg Precision | Neg Recall | Neg F1 | Micro F1 |")
    print(f"|---|---|---|---|---|---|---|---|")
    for r in summary_rows:
        print(f"| {r['pipeline']} "
              f"| {r['pos_precision']:.1%} | {r['pos_recall']:.1%} | {r['pos_f1']:.1%} "
              f"| {r['neg_precision']:.1%} | {r['neg_recall']:.1%} | {r['neg_f1']:.1%} "
              f"| {r['micro_f1']:.1%} |")


if __name__ == "__main__":
    main()
